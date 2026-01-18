"""
Proposal Document Generation Service

Generates refined proposal documents based on:
- Draft proposal content
- AI-generated section feedback
- User comments and guidance

The service filters to selected sections and uses Claude
to generate a refined proposal maintaining structure integrity.
"""

import json
import logging
import os
import re
import traceback
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

import boto3
from boto3.dynamodb.conditions import Attr
from docx import Document
from PyPDF2 import PdfReader

from app.shared.ai.bedrock_service import BedrockService
from app.tools.proposal_writer.proposal_document_generation.config import (
    PROPOSAL_DOCUMENT_GENERATION_SETTINGS,
)

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)


class ProposalDocumentGenerator:
    """
    Generates refined proposal documents using AI.

    Workflow:
    1. Load prompt template from DynamoDB
    2. Prepare context with draft, feedback, and user comments
    3. Filter to selected sections only
    4. Inject context into prompt
    5. Call Claude via Bedrock
    6. Parse and return refined proposal
    """

    def __init__(self):
        """Initialize Bedrock and DynamoDB clients."""
        self.bedrock = BedrockService()
        self.dynamodb = boto3.resource("dynamodb")
        self.s3 = boto3.client("s3")
        self.table_name = os.environ.get("TABLE_NAME", "igad-testing-main-table")
        self.bucket = os.environ.get("PROPOSALS_BUCKET")
        if not self.bucket:
            raise Exception("PROPOSALS_BUCKET environment variable not set")

    def generate_document(
        self,
        proposal_code: str,
        draft_proposal: str,
        section_feedback: List[Dict[str, Any]],
        user_comments: Dict[str, str],
        selected_sections: List[str],
    ) -> Dict[str, Any]:
        """
        Generate refined proposal document using AI.

        Args:
            proposal_code: Proposal identifier for logging
            draft_proposal: Full text of draft proposal
            section_feedback: AI feedback per section from draft_feedback_analysis
            user_comments: User-provided comments per section
            selected_sections: List of section titles to refine

        Returns:
            Dict with 'generated_proposal', 'sections', and 'metadata'

        Raises:
            ValueError: If prompt template not found
            Exception: If document generation fails
        """
        try:
            logger.info(f"📋 Generating refined proposal for: {proposal_code}")
            logger.info(f"   Selected sections: {len(selected_sections)}")
            logger.info(f"   User comments: {len(user_comments)}")

            start_time = datetime.utcnow()

            # Step 1: Load prompt template
            logger.info("📝 Loading prompt template...")
            prompt_parts = self._get_prompt_template()

            if not prompt_parts:
                raise ValueError("Prompt template not found in DynamoDB")

            # Step 2: Prepare context
            logger.info("🔄 Preparing context...")
            context = self._prepare_context(
                draft_proposal,
                section_feedback,
                user_comments,
                selected_sections,
            )

            # Step 3: Build final prompt
            user_prompt = self._inject_context(prompt_parts["user_prompt"], context)
            final_prompt = f"{user_prompt}\n\n{prompt_parts['output_format']}".strip()

            # Step 4: Call Bedrock
            logger.info("📡 Calling Bedrock (this may take 3-5 minutes)...")

            ai_response = self.bedrock.invoke_claude(
                system_prompt=prompt_parts["system_prompt"],
                user_prompt=final_prompt,
                max_tokens=PROPOSAL_DOCUMENT_GENERATION_SETTINGS.get(
                    "max_tokens", 16000
                ),
                temperature=PROPOSAL_DOCUMENT_GENERATION_SETTINGS.get(
                    "temperature", 0.2
                ),
                model_id=PROPOSAL_DOCUMENT_GENERATION_SETTINGS["model"],
            )

            elapsed = (datetime.utcnow() - start_time).total_seconds()
            logger.info(f"✅ Response received in {elapsed:.1f} seconds")

            # Step 5: Parse response
            logger.info("📊 Parsing response...")
            document = self._parse_response(ai_response)

            # Add metadata
            document["metadata"] = {
                "proposal_code": proposal_code,
                "sections_refined": len(selected_sections),
                "generation_time_seconds": elapsed,
                "generated_at": datetime.utcnow().isoformat(),
            }

            logger.info("✅ Refined proposal generated successfully")
            return document

        except Exception as e:
            logger.error(f"❌ Document generation failed: {str(e)}")
            traceback.print_exc()
            raise

    def _get_prompt_template(self) -> Optional[Dict[str, str]]:
        """
        Load document generation prompt from DynamoDB.

        Filters for:
        - is_active: true
        - section: "proposal_writer"
        - sub_section: "step-4"
        - categories: contains "Proposal Regeneration"

        Returns:
            Dict with 'system_prompt', 'user_prompt', 'output_format', or None
        """
        try:
            table = self.dynamodb.Table(self.table_name)
            filter_expr = (
                Attr("is_active").eq(True)
                & Attr("section").eq(PROPOSAL_DOCUMENT_GENERATION_SETTINGS["section"])
                & Attr("sub_section").eq(
                    PROPOSAL_DOCUMENT_GENERATION_SETTINGS["sub_section"]
                )
                & Attr("categories").contains(
                    PROPOSAL_DOCUMENT_GENERATION_SETTINGS["category"]
                )
            )

            # Handle DynamoDB pagination
            items = []
            response = table.scan(FilterExpression=filter_expr)
            items.extend(response.get("Items", []))

            while "LastEvaluatedKey" in response:
                response = table.scan(
                    FilterExpression=filter_expr,
                    ExclusiveStartKey=response["LastEvaluatedKey"],
                )
                items.extend(response.get("Items", []))

            if not items:
                logger.warning("⚠️  No prompts found in DynamoDB")
                return None

            prompt_item = items[0]
            logger.info(f"✅ Loaded prompt: {prompt_item.get('name', 'Unnamed')}")

            return {
                "system_prompt": prompt_item.get("system_prompt", ""),
                "user_prompt": prompt_item.get("user_prompt_template", ""),
                "output_format": prompt_item.get("output_format", ""),
            }

        except Exception as e:
            logger.error(f"❌ Error loading prompt: {str(e)}")
            return None

    def _prepare_context(
        self,
        draft_proposal: str,
        section_feedback: List[Dict[str, Any]],
        user_comments: Dict[str, str],
        selected_sections: List[str],
    ) -> Dict[str, str]:
        """
        Prepare context for prompt injection.

        Filters feedback to only selected sections.

        Args:
            draft_proposal: Full draft proposal text
            section_feedback: All section feedback
            user_comments: User comments keyed by section title
            selected_sections: List of selected section titles

        Returns:
            Dict with 'draft_proposal', 'section_feedback', 'user_comments'
        """
        # Filter feedback to selected sections only
        filtered_feedback = [
            f for f in section_feedback if f.get("section_title") in selected_sections
        ]

        logger.info(
            f"📊 Filtered to {len(filtered_feedback)} sections from {len(section_feedback)} total"
        )

        # Filter user comments to selected sections only
        filtered_comments = {
            k: v
            for k, v in user_comments.items()
            if k in selected_sections and v.strip()
        }

        logger.info(f"📝 Including {len(filtered_comments)} user comments")

        return {
            "draft_proposal": draft_proposal,
            "section_feedback": json.dumps(filtered_feedback, indent=2),
            "user_comments": json.dumps(filtered_comments, indent=2),
        }

    def _inject_context(self, template: str, context: Dict[str, Any]) -> str:
        """
        Inject context variables into prompt template.

        Supports placeholder formats:
        - {[DRAFT PROPOSAL]}
        - {[SECTION FEEDBACK]}
        - {[USER COMMENTS]}

        Args:
            template: Prompt template with placeholder markers
            context: Dict of context values

        Returns:
            Prompt with injected context
        """
        prompt = template
        replacements_made = 0

        for key, value in context.items():
            value_str = str(value)

            # Format 1: {[KEY]} (uppercase with spaces)
            key_upper = key.upper().replace("_", " ")
            placeholder_bracket = "{[" + key_upper + "]}"
            if placeholder_bracket in prompt:
                prompt = prompt.replace(placeholder_bracket, value_str)
                replacements_made += 1
                logger.info(f"✅ Replaced placeholder: {placeholder_bracket}")

            # Format 2: {[key]} (original key in bracket format)
            placeholder_bracket_original = "{[" + key + "]}"
            if placeholder_bracket_original in prompt:
                prompt = prompt.replace(placeholder_bracket_original, value_str)
                replacements_made += 1
                logger.info(f"✅ Replaced placeholder: {placeholder_bracket_original}")

            # Format 3: {{key}} (double brace format)
            placeholder_double_brace = f"{{{{{key}}}}}"
            if placeholder_double_brace in prompt:
                prompt = prompt.replace(placeholder_double_brace, value_str)
                replacements_made += 1
                logger.info(f"✅ Replaced placeholder: {placeholder_double_brace}")

        logger.info(f"🔄 Total placeholders replaced: {replacements_made}")

        # DEBUG: Check for any remaining unreplaced placeholders
        remaining_brackets = re.findall(r"\{\[[^\]]+\]\}", prompt)
        if remaining_brackets:
            logger.warning(f"⚠️ Unreplaced placeholders: {remaining_brackets[:5]}")

        return prompt

    def _parse_response(self, response: str) -> Dict[str, Any]:
        """
        Parse AI response into structured document.

        The AI returns the refined proposal as plain text/markdown.
        We extract sections by looking for markdown headers.

        Args:
            response: Raw response from Claude

        Returns:
            Dict with 'generated_proposal' and 'sections'
        """
        # The response should be the refined proposal document
        sections = self._extract_sections_from_text(response)

        return {
            "generated_proposal": response,
            "sections": sections,
        }

    def _extract_sections_from_text(self, text: str) -> Dict[str, str]:
        """
        Extract sections from markdown text.

        Looks for section headers:
        - ## Section Title (markdown h2)
        - # Section Title (markdown h1)

        Args:
            text: Markdown text

        Returns:
            Dict of section_title: content
        """
        sections = {}
        current_section = None
        current_content = []

        for line in text.split("\n"):
            stripped_line = line.strip()

            # Check for ## section headers (most common in proposals)
            if stripped_line.startswith("## "):
                if current_section:
                    sections[current_section] = "\n".join(current_content).strip()
                current_section = stripped_line[3:].strip()
                current_content = []
            # Check for # section headers
            elif stripped_line.startswith("# ") and not stripped_line.startswith("## "):
                if current_section:
                    sections[current_section] = "\n".join(current_content).strip()
                current_section = stripped_line[2:].strip()
                current_content = []
            elif current_section:
                current_content.append(line)

        # Don't forget the last section
        if current_section:
            sections[current_section] = "\n".join(current_content).strip()

        logger.info(f"📊 Extracted {len(sections)} sections from text")
        for title in list(sections.keys())[:5]:
            logger.info(f"   ✓ {title}")

        return sections


# ==================== SERVICE INSTANCE ====================

# Global instance for document generation
proposal_document_generator = ProposalDocumentGenerator()
